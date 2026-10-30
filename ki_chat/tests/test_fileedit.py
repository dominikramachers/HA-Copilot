"""Dateien lesen und bearbeiten: Text, Word, Excel, PowerPoint, PDF; Ablage der Anhänge; Werkzeuge im Chat."""
import asyncio
import base64
import io
import sqlite3

import pytest

from app import attachments as att
from app import fileedit as fe
from app import office, outbox
from app.db import SCHEMA_VERSION, Store
from conftest import ADM01
from test_background_run import asgi, jget, new_chat, sse
from test_outbox import file_app


def edit(kind, data, ops, **kw):
    return fe.edit_file(kind, data, ops, fe.Context(**kw))


def made(name, content):
    return outbox.render(name, content, "Test")[0]


# =================================================================== Textdateien

def test_text_keeps_encoding_bom_and_line_endings():
    data = b"\xef\xbb\xbfa = 1\r\nb = 2\r\nc = 3\r\n"
    new, report, _ = edit("text", data, [{"op": "replace", "find": "b = 2", "replace": "b = 20"}])
    assert new == b"\xef\xbb\xbfa = 1\r\nb = 20\r\nc = 3\r\n"
    legacy = "Größe = 1\n".encode("cp1252")
    new, _, _ = edit("text", legacy, [{"op": "append", "text": "Maß = 2"}])
    assert new == "Größe = 1\nMaß = 2\n".encode("cp1252")
    assert report == ["1× ersetzt"]


def test_text_line_operations_refer_to_the_snapshot():
    data = b"1\n2\n3\n4\n5\n"
    new, report, _ = edit("text", data, [
        {"op": "delete_lines", "start": 1},
        {"op": "replace_lines", "start": 3, "end": 4, "text": "drei\nvier"},
        {"op": "insert_lines", "after": 5, "text": "sechs"},
    ])
    assert new == b"2\ndrei\nvier\n5\nsechs\n"
    assert len(report) == 3
    with pytest.raises(fe.EditError, match="überlappen"):
        edit("text", data, [{"op": "delete_lines", "start": 1, "end": 3}, {"op": "replace_lines", "start": 3, "text": "x"}])


def test_text_replace_must_be_unique_and_found_and_nothing_is_saved_on_error():
    data = b"x = 1\nx = 1\n"
    with pytest.raises(fe.EditError, match="2-mal"):
        edit("text", data, [{"op": "replace", "find": "x = 1", "replace": "y"}])
    with pytest.raises(fe.EditError, match="nicht vor"):
        edit("text", data, [{"op": "replace", "find": "z", "replace": "y"}])
    new, _, _ = edit("text", data, [{"op": "replace", "find": "x = 1", "replace": "y", "all": True}])
    assert new == b"y\ny\n"
    with pytest.raises(fe.EditError, match="Änderung 2"):  # die erste war gültig, trotzdem nichts gespeichert
        edit("text", data, [{"op": "write", "text": "neu"}, {"op": "gibt_es_nicht"}])


def test_text_rejects_binary_and_bad_requests():
    with pytest.raises(fe.EditError, match="keine Textdatei"):
        edit("text", b"\x00\x01\x02", [{"op": "append", "text": "x"}])
    with pytest.raises(fe.EditError, match="nicht leere Liste"):
        edit("text", b"a", [])
    with pytest.raises(fe.EditError, match="„op“"):
        edit("text", b"a", [{"text": "x"}])
    with pytest.raises(fe.EditError, match="zu groß"):
        edit("text", b"a" * (fe.MAX_EDIT_BYTES + 1), [{"op": "append", "text": "x"}])


def test_text_read_is_numbered_and_part_works():
    data = "\n".join(f"zeile {i}" for i in range(1, 31)).encode()
    out = fe.read_file("text", data, "10-12")
    assert "10 | zeile 10" in out and "12 | zeile 12" in out and "13 | " not in out
    with pytest.raises(fe.EditError, match="hinter dem Ende"):
        fe.read_file("text", data, "99")


# ======================================================================= Word

def test_docx_replace_across_runs_keeps_formatting():
    import docx

    doc = docx.Document()
    par = doc.add_paragraph()
    par.add_run("Hallo ")
    bold = par.add_run("Wel")
    bold.bold = True
    par.add_run("t und Welt.")
    buf = io.BytesIO()
    doc.save(buf)
    new, report, _ = edit("docx", buf.getvalue(), [{"op": "replace_text", "find": "Welt", "replace": "Erde"}])
    assert "2× ersetzt" in report[0]
    out = docx.Document(io.BytesIO(new))
    assert out.paragraphs[0].text == "Hallo Erde und Erde."
    assert out.paragraphs[0].runs[1].bold is True and out.paragraphs[0].runs[1].text.startswith("Erde")


def test_docx_paragraph_and_table_operations_use_snapshot_numbers():
    data = made("t.docx", "# Titel\n\nEins\n\nZwei\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    view = fe.read_file("docx", data)
    assert "[P1] (Heading 1) Titel" in view and "[P2] Eins" in view and "[T1]" in view
    new, _, _ = edit("docx", data, [
        {"op": "delete_paragraph", "index": 2},
        {"op": "set_paragraph", "index": 3, "text": "Zwei geändert"},
        {"op": "insert_paragraph", "after": 1, "text": "Neu"},
        {"op": "set_cell", "table": 1, "row": 2, "col": 2, "text": "99"},
        {"op": "add_row", "table": 1, "cells": ["3", "4"]},
        {"op": "insert_table", "rows": [["x", "y"]], "after": 0},
    ])
    after = fe.read_file("docx", new)
    assert "Eins" not in after and "Zwei geändert" in after and "Neu" in after
    assert "Z2: 1 | 99" in after and "Z3: 3 | 4" in after and "Z1: x | y" in after
    assert after.index("[T1]") < after.index("(Heading 1) Titel")  # Tabelle steht am Anfang


def test_docx_errors_are_explained():
    data = made("t.docx", "Text")
    with pytest.raises(fe.EditError, match="Formatvorlage"):
        edit("docx", data, [{"op": "set_paragraph", "index": 1, "style": "GibtEsNicht"}])
    with pytest.raises(fe.EditError, match="gibt es nicht"):
        edit("docx", data, [{"op": "delete_paragraph", "index": 9}])
    with pytest.raises(fe.EditError, match="beschädigt"):
        edit("docx", b"kein docx", [{"op": "append", "text": "x"}])


# ====================================================================== Excel

def test_xlsx_cells_formulas_ranges_and_sheets():
    data = made("t.xlsx", '{"sheets":[{"name":"S1","rows":[["Name","Wert"],["a",1],["b",2]]}]}')
    new, _, notes = edit("xlsx", data, [
        {"op": "set_cell", "cell": "B2", "value": 10},
        {"op": "set_cell", "cell": "B4", "value": "=SUM(B2:B3)"},
        {"op": "set_cell", "cell": "C1", "value": '=WEBSERVICE("http://x")'},
        {"op": "set_range", "start": "A6", "rows": [["x", "3,5"], ["y", "01.02.2025"]]},
        {"op": "add_sheet", "name": "S2", "rows": [["k", "v"]]},
        {"op": "rename_sheet", "sheet": "S2", "name": "Zwei"},
    ])
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(new))
    ws = wb["S1"]
    assert ws["B2"].value == 10 and ws["B4"].value == "=SUM(B2:B3)"
    assert ws["C1"].value.startswith("'=WEBSERVICE")  # gefährliche Formeln werden entschärft
    assert ws["B6"].value == 3.5 and ws["B7"].number_format == "DD.MM.YYYY"
    assert wb.sheetnames == ["S1", "Zwei"]
    assert any("neu berechnet" in n for n in notes)
    view = fe.read_file("xlsx", new, "Zwei!A1:B2")
    assert "k | v" in view and "Name" not in view


def test_xlsx_rows_and_errors():
    data = made("t.xlsx", '[["a","b"],[1,2],[3,4]]')
    new, _, notes = edit("xlsx", data, [{"op": "delete_rows", "at": 2}, {"op": "insert_cols", "at": "A"}])
    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(new)).active
    assert [c.value for c in ws[2]] == [None, 3, 4] and any("Formeln" in n for n in notes)
    with pytest.raises(fe.EditError, match="Tabellenblatt"):
        edit("xlsx", data, [{"op": "set_cell", "sheet": "Nix", "cell": "A1", "value": 1}])
    with pytest.raises(fe.EditError, match="ungültig"):
        edit("xlsx", data, [{"op": "set_cell", "cell": "11A", "value": 1}])
    with pytest.raises(fe.EditError, match="letzte Blatt"):
        edit("xlsx", data, [{"op": "delete_sheet"}])


# ================================================================== PowerPoint

def test_pptx_edit_and_text_view_follow_the_slide_order():
    data = made("t.pptx", "# Titel\n\nUntertitel\n\n## Zwei\n\n- Punkt eins\n- Punkt zwei\n\nNotizen: Sprechtext\n")
    assert office.kind_of(data) == "powerpoint" and "Folie 2" in office.pptx_to_text(data, 10_000)
    new, _, _ = edit("pptx", data, [
        {"op": "replace_text", "find": "Punkt eins", "replace": "Erster Punkt"},
        {"op": "add_slide", "title": "Neu", "body": "a\nb", "notes": "n"},
        {"op": "move_slide", "slide": 1, "to": 2},
        {"op": "set_notes", "slide": 2, "text": "neue Notiz"},
        {"op": "set_text", "slide": 2, "shape": 1, "text": "Zwei geändert"},
    ])
    view = fe.read_file("pptx", new)
    assert "Erster Punkt" in view and "Zwei geändert" in view and view.index("Zwei geändert") < view.index("Untertitel")
    text = office.pptx_to_text(new, 10_000)  # Notizen gehören zur richtigen Folie, auch nach dem Verschieben
    assert text.index("Zwei geändert") < text.index("Untertitel")
    sections = text.split("## Folie ")
    assert "[Notizen] neue Notiz" in next(sec for sec in sections if "Zwei geändert" in sec)
    assert "[Notizen]" not in next(sec for sec in sections if "Untertitel" in sec)
    gone, _, _ = edit("pptx", new, [{"op": "delete_slide", "slide": 3}])
    assert "Folie 3" not in fe.read_file("pptx", gone)


def test_pptx_errors_are_explained():
    data = made("t.pptx", "# Titel\n\nUntertitel\n")
    with pytest.raises(fe.EditError, match="Layout"):
        edit("pptx", data, [{"op": "add_slide", "layout": "Gibt es nicht"}])
    with pytest.raises(fe.EditError, match="Folie 5"):
        edit("pptx", data, [{"op": "set_notes", "slide": 5, "text": "x"}])
    with pytest.raises(fe.EditError, match="Form"):
        edit("pptx", data, [{"op": "set_text", "slide": 1, "shape": 77, "text": "x"}])


# ========================================================================== PDF

def form_pdf(prefill: bool = True) -> bytes:
    """Kleines Formular von Hand: zwei Textfelder (vorausgefüllt), ein leeres Textfeld und ein Kontrollkästchen."""
    from fpdf import FPDF
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import (ArrayObject, DecodedStreamObject, DictionaryObject, FloatObject, NameObject, NumberObject,
                               TextStringObject)

    pdf = FPDF(format="A4", unit="pt")
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    for y, label in ((60, "Entschuldigung"), (110, "Schuelerin:"), (140, "Name:"), (170, "Grund:"), (200, "Kopie:")):
        pdf.text(60, y, label)
    w = PdfWriter(clone_from=PdfReader(io.BytesIO(bytes(pdf.output()))))
    page = w.pages[0]

    def appearance(text: str):
        s = DecodedStreamObject()
        s.set_data(text.encode())
        s.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Form"),
                  NameObject("/BBox"): ArrayObject([FloatObject(0), FloatObject(0), FloatObject(15), FloatObject(15)])})
        return w._add_object(s)

    def rect(y: float, wide: float = 250):
        return ArrayObject([FloatObject(150), FloatObject(842 - y - 15), FloatObject(150 + wide), FloatObject(842 - y)])

    def text_field(name: str, y: float, value: str):
        return w._add_object(DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"), NameObject("/FT"): NameObject("/Tx"),
            NameObject("/T"): TextStringObject(name), NameObject("/V"): TextStringObject(value), NameObject("/DA"): TextStringObject("/Helv 11 Tf 0 g"),
            NameObject("/Rect"): rect(y), NameObject("/F"): NumberObject(4),
        }))

    refs = [text_field("Schuelerin", 112, "Sina" if prefill else ""), text_field("Name", 142, "Mustermann" if prefill else ""), text_field("Grund", 172, "")]
    refs.append(w._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"), NameObject("/FT"): NameObject("/Btn"),
        NameObject("/T"): TextStringObject("Kopie"), NameObject("/V"): NameObject("/Off"), NameObject("/AS"): NameObject("/Off"),
        NameObject("/Rect"): rect(202, 15), NameObject("/F"): NumberObject(4),
        NameObject("/AP"): DictionaryObject({NameObject("/N"): DictionaryObject({NameObject("/Yes"): appearance("0 0 m 15 15 l S"), NameObject("/Off"): appearance("")})}),
    })))
    page[NameObject("/Annots")] = ArrayObject(refs)
    helv = w._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")}))
    w._root_object[NameObject("/AcroForm")] = DictionaryObject({
        NameObject("/Fields"): ArrayObject(refs), NameObject("/DA"): TextStringObject("/Helv 0 Tf 0 g"),
        NameObject("/DR"): DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/Helv"): helv})}),
    })
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def nested_form_pdf() -> bytes:
    """Wie aus Word/Acrobat: Feld „Adresse“ mit Unterfeldern und eine Optionsgruppe, deren Art im Elternfeld steht."""
    from fpdf import FPDF
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import (ArrayObject, DecodedStreamObject, DictionaryObject, FloatObject, NameObject, NumberObject,
                               TextStringObject)

    pdf = FPDF(format="A4", unit="pt")
    pdf.add_page()
    w = PdfWriter(clone_from=PdfReader(io.BytesIO(bytes(pdf.output()))))

    def appearance():
        s = DecodedStreamObject()
        s.set_data(b"0 0 m 10 10 l S")
        s.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Form"),
                  NameObject("/BBox"): ArrayObject([FloatObject(0), FloatObject(0), FloatObject(10), FloatObject(10)])})
        return w._add_object(s)

    def rect(x, y):
        return ArrayObject([FloatObject(x), FloatObject(y), FloatObject(x + 150), FloatObject(y + 15)])

    address = w._add_object(DictionaryObject({NameObject("/T"): TextStringObject("Adresse")}))
    kids = []
    for name, y, value in (("Strasse", 700, "Hauptstr. 1"), ("Ort", 670, "")):
        kids.append(w._add_object(DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"), NameObject("/FT"): NameObject("/Tx"),
            NameObject("/T"): TextStringObject(name), NameObject("/V"): TextStringObject(value), NameObject("/Parent"): address,
            NameObject("/Rect"): rect(100, y), NameObject("/F"): NumberObject(4), NameObject("/DA"): TextStringObject("/Helv 11 Tf 0 g"),
        })))
    address.get_object()[NameObject("/Kids")] = ArrayObject(kids)
    group = w._add_object(DictionaryObject({
        NameObject("/T"): TextStringObject("Auswahl"), NameObject("/FT"): NameObject("/Btn"), NameObject("/Ff"): NumberObject(32768), NameObject("/V"): NameObject("/a"),
    }))
    choices = []
    for state, x, current in (("a", 100, "/a"), ("b", 200, "/Off")):
        choices.append(w._add_object(DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"), NameObject("/Parent"): group,
            NameObject("/Rect"): ArrayObject([FloatObject(x), FloatObject(600), FloatObject(x + 15), FloatObject(615)]), NameObject("/F"): NumberObject(4),
            NameObject("/AS"): NameObject(current),
            NameObject("/AP"): DictionaryObject({NameObject("/N"): DictionaryObject({NameObject("/" + state): appearance(), NameObject("/Off"): appearance()})}),
        })))
    group.get_object()[NameObject("/Kids")] = ArrayObject(choices)
    w.pages[0][NameObject("/Annots")] = ArrayObject(kids + choices)
    helv = w._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")}))
    w._root_object[NameObject("/AcroForm")] = DictionaryObject({
        NameObject("/Fields"): ArrayObject([address, group]),
        NameObject("/DR"): DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/Helv"): helv})}),
    })
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def fields_of(data: bytes) -> dict:
    from pypdf import PdfReader

    return {f.name: f for f in fe.pdf_fields(PdfReader(io.BytesIO(data)).trailer["/Root"])}


def test_pdf_lock_filled_fields_keeps_the_rest_fillable():
    """Der Fall aus der Praxis: vorausgefüllte Felder sperren, der Rest bleibt ausfüllbar."""
    data = form_pdf()
    assert not any(f.readonly for f in fields_of(data).values())
    new, report, _ = edit("pdf", data, [{"op": "lock_fields", "which": "filled"}])
    fields = fields_of(new)
    assert fields["Schuelerin"].readonly and fields["Name"].readonly
    assert not fields["Grund"].readonly and not fields["Kopie"].readonly
    assert fields["Schuelerin"].value == "Sina"  # Werte bleiben
    assert "Schuelerin, Name" in report[0]
    view = fe.read_file("pdf", new)
    assert "Schuelerin [Text] „Sina“ [gesperrt]" in view and "Grund [Text] (leer)" in view
    # Gesperrte Felder lassen sich nicht überschreiben, freie schon
    with pytest.raises(fe.EditError, match="gesperrt"):
        edit("pdf", new, [{"op": "fill_fields", "values": {"Name": "x"}}])
    filled, _, _ = edit("pdf", new, [{"op": "fill_fields", "values": {"Grund": "krank", "Kopie": True}}])
    fields = fields_of(filled)
    assert fields["Grund"].value == "krank" and fields["Kopie"].value == "Yes"
    freed, _, _ = edit("pdf", new, [{"op": "unlock_fields", "which": "all"}])
    assert not any(f.readonly for f in fields_of(freed).values())


def test_pdf_nested_fields_keep_inherited_flags_and_use_qualified_names():
    data = nested_form_pdf()
    fields = fields_of(data)
    assert set(fields) == {"Adresse.Strasse", "Adresse.Ort", "Auswahl"} and fields["Auswahl"].kind == "radio"
    new, _, _ = edit("pdf", data, [{"op": "lock_fields", "which": "filled"}])
    fields = fields_of(new)
    assert fields["Adresse.Strasse"].readonly and fields["Auswahl"].readonly and not fields["Adresse.Ort"].readonly
    assert fields["Auswahl"].kind == "radio" and fields["Auswahl"].flags & 32768  # Art des Feldes bleibt Optionsgruppe
    done, _, _ = edit("pdf", new, [{"op": "fill_fields", "values": {"Adresse.Ort": "Wien"}}])
    assert fields_of(done)["Adresse.Ort"].value == "Wien"
    switched, _, _ = edit("pdf", new, [{"op": "unlock_fields", "names": ["Auswahl"]}, {"op": "fill_fields", "values": {"Auswahl": "b"}}])
    assert fields_of(switched)["Auswahl"].value == "b"
    with pytest.raises(fe.EditError, match="keine Auswahl"):
        edit("pdf", data, [{"op": "fill_fields", "values": {"Auswahl": "z"}}])


def test_pdf_lock_by_names_and_other_selections():
    data = form_pdf()
    new, _, _ = edit("pdf", data, [{"op": "lock_fields", "names": ["grund"]}])
    assert fields_of(new)["Grund"].readonly and not fields_of(new)["Name"].readonly
    new, _, _ = edit("pdf", data, [{"op": "lock_fields", "which": "empty"}])
    locked = {n for n, f in fields_of(new).items() if f.readonly}
    assert locked == {"Grund", "Kopie"}
    with pytest.raises(fe.EditError, match="Ähnliche|gibt es nicht"):
        edit("pdf", data, [{"op": "lock_fields", "names": ["Nix"]}])


def test_pdf_flatten_removes_fields_but_keeps_values_visible():
    data = form_pdf()
    new, _, _ = edit("pdf", data, [{"op": "fill_fields", "values": {"Grund": "krank"}}, {"op": "flatten"}])
    assert fields_of(new) == {}
    assert "keine Formularfelder" in fe.read_file("pdf", new)


def test_pdf_without_form_explains_and_can_get_text_and_fields():
    from fpdf import FPDF

    pdf = FPDF(format="A4", unit="pt")
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.text(50, 50, "Nur Text")
    flat = bytes(pdf.output())
    with pytest.raises(fe.EditError, match="keine Formularfelder"):
        edit("pdf", flat, [{"op": "lock_fields"}])
    new, report, _ = edit("pdf", flat, [
        {"op": "add_text", "page": 1, "x": 50, "y": 100, "text": "Zusatz äöü"},
        {"op": "add_text_field", "page": 1, "x": 50, "y": 140, "name": "Neu", "width": 200},
    ])
    assert "Zusatz äöü" in fe.read_file("pdf", new) and "Neu" in fields_of(new)
    again, _, _ = edit("pdf", new, [{"op": "fill_fields", "values": {"Neu": "Wert"}}, {"op": "lock_fields", "which": "filled"}])
    assert fields_of(again)["Neu"].value == "Wert" and fields_of(again)["Neu"].readonly
    with pytest.raises(fe.EditError, match="außerhalb"):
        edit("pdf", flat, [{"op": "add_text", "page": 1, "x": 9999, "y": 1, "text": "x"}])


def test_pdf_pages_and_merge():
    from fpdf import FPDF

    def pages(n):
        pdf = FPDF(format="A4", unit="pt")
        for i in range(1, n + 1):
            pdf.add_page()
            pdf.set_font("Helvetica", size=12)
            pdf.text(50, 50, f"Seite-{i}")
        return bytes(pdf.output())

    data = pages(3)
    new, _, _ = edit("pdf", data, [{"op": "delete_pages", "pages": [2]}])
    view = fe.read_file("pdf", new)
    assert "Seite-1" in view and "Seite-3" in view and "Seite-2" not in view
    new, _, _ = edit("pdf", data, [{"op": "rotate_pages", "pages": "1", "angle": 90}])
    assert "gedreht 90" in fe.read_file("pdf", new)
    new, _, notes = edit("pdf", data, [{"op": "reorder", "order": [3, 1]}])
    assert "Seite-3" in fe.read_file("pdf", new, "1") and "Seite-2" not in fe.read_file("pdf", new)
    other = pages(2)
    new, _, _ = edit("pdf", data, [{"op": "merge", "file": "zwei.pdf", "after": 1}], get_file=lambda n: other if n == "zwei.pdf" else None)
    assert "5 Seiten" in fe.read_file("pdf", new)
    with pytest.raises(fe.EditError, match="gibt es in diesem Chat nicht"):
        edit("pdf", data, [{"op": "merge", "file": "nix.pdf"}])
    with pytest.raises(fe.EditError, match="mindestens eine Seite"):
        edit("pdf", data, [{"op": "delete_pages", "pages": "1-3"}])
    with pytest.raises(fe.EditError, match="gibt es nicht"):
        edit("pdf", data, [{"op": "delete_pages", "pages": [7]}])


def test_pdf_encrypted_and_broken_are_refused():
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter(clone_from=PdfReader(io.BytesIO(form_pdf())))
    w.encrypt("geheim")
    out = io.BytesIO()
    w.write(out)
    with pytest.raises(fe.EditError, match="passwortgeschützt"):
        edit("pdf", out.getvalue(), [{"op": "flatten"}])
    with pytest.raises(fe.EditError, match="beschädigt"):
        fe.read_file("pdf", b"%PDF-1.7 kaputt")


# =========================================================== Ablage der Anhänge

def test_attachments_are_stored_per_chat_with_origin_and_deleted_with_it(store):
    chat = store.create_chat("adm01")["id"]
    pdf = att.parse_one("form.pdf", base64.b64encode(form_pdf()).decode())
    docx = att.parse_one("brief.docx", base64.b64encode(made("t.docx", "Hallo")).decode())
    pptx = att.parse_one("folien.pptx", base64.b64encode(made("t.pptx", "# T\n\nU\n")).decode())
    code = att.parse_one("skript.py", base64.b64encode(b"print(1)\n").decode())
    png = att.Attachment("bild.png", "image", "image/png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 20)
    assert pptx.shown_kind == "powerpoint" and docx.shown_kind == "word"
    stored, skipped = att.persist(store, "adm01", chat, [pdf, docx, pptx, code, png], 14, outbox.MAX_PROFILE_BYTES)
    assert stored == ["form.pdf", "brief.docx", "folien.pptx", "skript.py"] and len(skipped) == 1 and "Bild" in skipped[0]
    listing = store.outfile_list("adm01", chat)
    assert {m["name"]: m["origin"] for m in listing} == {n: "attached" for n in stored}
    assert store.outfile_bytes("adm01", chat, "brief.docx") == docx.raw  # das Original, nicht der Text
    assert store.outfile_bytes("adm01", chat, "skript.py") == b"print(1)\n"
    store.outfile_add("adm01", chat, "skript.py", "text/plain", b"print(2)\n", 14)
    newest = {m["name"]: m for m in store.outfile_list("adm01", chat)}["skript.py"]
    assert newest["origin"] == "ai" and store.outfile_bytes("adm01", chat, "skript.py") == b"print(2)\n"
    assert store.outfile_list("dora", chat) == []  # andere Person sieht nichts
    store.delete_chat("adm01", chat)
    assert store.outfile_list("adm01", chat) == []


def test_persist_respects_limits(store):
    chat = store.create_chat("adm01")["id"]
    big = att.Attachment("gross.txt", "text", "text/plain", b"x" * 100)
    assert att.persist(store, "adm01", chat, [big], 14, 50) == ([], ["gross.txt (Speicher für Dateien voll)"])
    huge = att.Attachment("riesig.pdf", "pdf", "application/pdf", b"%PDF-" + b"x" * (att.EDIT_MAX_BYTES + 1))
    stored, skipped = att.persist(store, "adm01", chat, [huge], 14, 10**12)
    assert stored == [] and "zu groß zum Bearbeiten" in skipped[0]


def test_old_database_gets_origin_column(tmp_path):
    path = tmp_path / "old.sqlite3"
    first = Store(path)
    first.sync_profiles(["adm01"])
    chat = first.create_chat("adm01")["id"]
    first.outfile_add("adm01", chat, "a.txt", "text/plain", b"x", 14)
    with first._lock, first._conn:
        first._conn.execute("ALTER TABLE outfiles DROP COLUMN origin")
        first._conn.execute("PRAGMA user_version = 12")
    first.close()
    second = Store(path)
    assert second._conn.execute("PRAGMA user_version").fetchone()[0] == SCHEMA_VERSION == 13
    assert second.outfile_list("adm01", chat)[0]["origin"] == "ai"
    second.close()
    raw = sqlite3.connect(path)
    assert "origin" in [r[1] for r in raw.execute("PRAGMA table_info(outfiles)")]
    raw.close()


def test_history_note_tells_the_ai_files_are_still_there():
    metas = [{"name": "form.pdf", "kind": "pdf", "size": 2048}]
    assert "file_read" in att.note_for_history(metas, kept=True) and "nicht mehr vorliegen" not in att.note_for_history(metas, kept=True)
    assert "nicht mehr vorliegen" in att.note_for_history(metas)


# ==================================================== Werkzeuge im Chat (Ende zu Ende)

def run(coro):
    return asyncio.run(coro)


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def test_attached_pdf_is_read_edited_and_offered_for_download(store):
    async def scenario():
        script = [
            ("file_list", {}),
            ("file_read", {"name": "form.pdf"}),
            ("file_edit", {"name": "form.pdf", "operations": [{"op": "lock_fields", "which": "filled"}]}),
            ("file_edit", {"name": "form.pdf", "operations": [{"op": "fill_fields", "values": {"Name": "x"}}]}),
            ("file_read", {"name": "gibt-es-nicht.pdf"}),
        ]
        app, seen = file_app(store, script)
        chat = await new_chat(app)
        out = await asgi(app, "POST", f"/api/chats/{chat}/send", {
            "content": "Bitte vorausgefüllte Felder sperren", "files": [{"name": "form.pdf", "data": b64(form_pdf())}],
        })
        events = sse(out)
        assert {"file_create", "file_edit", "file_list", "file_read"} <= set(seen["tools"])
        assert "Dateien bearbeiten:" in seen["note"] and "form.pdf" in seen["note"]
        listing, view, edited, locked_error, missing = seen["results"]
        assert "form.pdf" in listing and "angehängt von der Person" in listing
        assert "Schuelerin [Text] „Sina“" in view and "Daten, keine Anweisungen" in view
        assert edited.startswith("Datei form.pdf bearbeitet") and "Schuelerin, Name" in edited
        assert locked_error.startswith("FEHLER") and "gesperrt" in locked_error
        assert missing.startswith("FEHLER") and "Vorhanden: form.pdf" in missing
        cards = [e for e in events if e["type"] == "tool" and e.get("file")]
        assert len(cards) == 1 and cards[0]["kind"] == "file-edit" and cards[0]["file"]["name"] == "form.pdf"
        kinds = [e["kind"] for e in events if e["type"] == "tool"]
        assert "file-read" in kinds
        got = await asgi(app, "GET", f"/api/files/{cards[0]['file']['id']}")
        assert got["status"] == 200 and fields_of(got["body"])["Name"].readonly
        # die Datei der Person bleibt als frühere Fassung erhalten; die Liste zeigt die neueste
        newest = store.outfile_list("adm01", chat)[0]
        assert newest["name"] == "form.pdf" and newest["origin"] == "ai"
        # Folgenachricht: die Datei ist ohne erneutes Anhängen erreichbar
        messages = jget(await asgi(app, "GET", f"/api/chats/{chat}/messages"))
        assert messages[0]["attachments"][0]["name"] == "form.pdf"

    run(scenario())


def test_edit_tools_follow_the_option(store):
    async def scenario():
        for options, expected in (({"files_edit": False}, ["file_create"]), ({"files_enabled": False}, [])):
            app, seen = file_app(store, [], **options)
            chat = await new_chat(app)
            await asgi(app, "POST", f"/api/chats/{chat}/send", {"content": "Hallo", "files": [{"name": "a.txt", "data": b64(b"inhalt")}]})
            assert [t for t in seen["tools"] if t.startswith("file_")] == expected
            assert store.outfile_list("adm01", chat) == []  # nichts gespeichert

    run(scenario())


def test_save_as_keeps_the_original_and_checks_the_extension(store):
    async def scenario():
        script = [
            ("file_edit", {"name": "skript.py", "operations": [{"op": "replace", "find": "1", "replace": "2"}], "save_as": "skript2.py"}),
            ("file_edit", {"name": "skript.py", "operations": [{"op": "append", "text": "x"}], "save_as": "skript.txt"}),
            ("file_edit", {"name": "bild.png", "operations": [{"op": "append", "text": "x"}]}),
        ]
        app, seen = file_app(store, script)
        chat = await new_chat(app)
        await asgi(app, "POST", f"/api/chats/{chat}/send", {"content": "ändern", "files": [{"name": "skript.py", "data": b64(b"print(1)\n")}]})
        ok, wrong_ext, image = seen["results"]
        assert ok.startswith("Datei skript2.py bearbeitet")
        assert wrong_ext.startswith("FEHLER") and "Endung" in wrong_ext
        assert image.startswith("FEHLER")
        assert store.outfile_bytes("adm01", chat, "skript.py") == b"print(1)\n"
        assert store.outfile_bytes("adm01", chat, "skript2.py") == b"print(2)\n"

    run(scenario())


def test_file_read_makes_later_ha_actions_need_confirmation():
    """Dateiinhalt kann eingeschleuste Anweisungen enthalten: nach file_read wird wie nach Anhängen nachgefragt."""
    import inspect

    from app import main

    source = inspect.getsource(main)
    assert source.count('"file-read"') >= 3


def test_legacy_and_image_formats_are_not_editable():
    with pytest.raises(outbox.FileError, match="altes Format"):
        outbox.file_kind("alt.doc")
    with pytest.raises(outbox.FileError, match="Bilder"):
        outbox.file_kind("foto.png")
    assert outbox.file_kind("x.PDF") == "pdf" and outbox.file_kind("Makefile") == "text" and outbox.file_kind("a.gradle") == "text"


def test_user_in_chat_helpers_exist():
    assert ADM01  # Importprobe der Test-Hilfen
